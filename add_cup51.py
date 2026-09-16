"""Add Cup 51 (S4 Round 1) to Petite Cups 51-55.xlsx in cols 1-4.

First cup of Season 4, and the first in a new workbook — Petite Cups 46-50.xlsx
filled up at Cup 50.

BEFORE RUNNING:
  1. Set EXCLUDE to the mapper(s) who were in the lobby but not competing.
     Petite uses TWO maps, so there are usually two mappers. If a mapper sat the
     cup out entirely, they won't be in the log and need no exclusion.
  2. Set MAPS to the row-3 string once the map names are confirmed.
  3. Confirm cup logs/petite_51_reconstructed.json exists (run parse_petite_log.py 51).

AFTER RUNNING, in petite_ranking.py:
  FULL_LOBBY_REPLACEMENTS['Season 4'] = {'Round 1': ('Petite Cups 51-55.xlsx', 'Petite Cup 51')}
  EVENT_DATES['Season 4'] = {'Round 1': '<YYYY-MM-DD>'}
Both are required — the xlsx alone puts data in the file but not in the rankings.
"""
import json
import os
import openpyxl

base = os.path.dirname(os.path.abspath(__file__))
xlsx_path = os.path.join(base, 'Petite Cups 51-55.xlsx')
json_path = os.path.join(base, 'cup logs', 'petite_51_reconstructed.json')

EXCLUDE = set()  # <-- mapper(s) in the lobby who weren't competing
MAPS = 'Maps: TODO'  # <-- row 3 map string

with open(json_path, 'r', encoding='utf-8') as f:
    raw = json.load(f)

filtered = [r for r in raw if r['name'] not in EXCLUDE]

new_list = []
prev_pos = None
new_pos = 0
seen = 0
for entry in filtered:
    seen += 1
    if entry['pos'] != prev_pos:
        new_pos = seen
        prev_pos = entry['pos']
    new_list.append({**entry, 'new_pos': new_pos})

print(f'Total entries after mapper exclusion: {len(new_list)}')
for e in new_list:
    t = f"{e['time']:.5f}" if e['time'] is not None else 'DNF'
    safe = e['name'].encode('ascii', 'replace').decode('ascii')
    print(f"  {e['new_pos']:3} {safe:38} {t:>10}  (R{e['round'] or '-'})")

wb = openpyxl.load_workbook(xlsx_path)
ws = wb.active

COL = 1  # cup 51 at 1, 52 at 7, 53 at 13, 54 at 19, 55 at 25 (6-col groups)
ws.cell(row=2, column=COL, value='Petite Cup 51')
ws.cell(row=3, column=COL, value=MAPS)
ws.cell(row=5, column=COL,   value='Position')
ws.cell(row=5, column=COL+1, value='Name')
ws.cell(row=5, column=COL+2, value='Elim Time')
ws.cell(row=5, column=COL+3, value='Elim Round')

for r in range(6, 60):
    for c in range(COL, COL+4):
        ws.cell(row=r, column=c).value = None

for i, e in enumerate(new_list):
    r = 6 + i
    ws.cell(row=r, column=COL,   value=e['new_pos'])
    ws.cell(row=r, column=COL+1, value=e['name'])
    if e['time'] is not None:
        ws.cell(row=r, column=COL+2, value=round(e['time'], 5))
    else:
        ws.cell(row=r, column=COL+2, value='DNF')
    if e['round'] is not None and e.get('note') != 'winner':
        ws.cell(row=r, column=COL+3, value=e['round'])

wb.save(xlsx_path)
print(f'\nWrote Cup 51 -> {xlsx_path}')
