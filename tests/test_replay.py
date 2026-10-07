"""new_petite.py rebuilds Cups 51-53 exactly, and its guards refuse what they should.

Everything runs in a scratch copy with PETITE_SKIP_EXTERNAL=1; the real repo,
the cross-comp data and port 9001 are never touched.
"""
import json
import shutil

import openpyxl
import pytest

from conftest import REPO, run_script

LOGS = REPO / 'cup logs'
BOOK = 'Petite Cups 51-55.xlsx'

CUPS = {
    51: ['--community', '50', '--date', '2026-09-16',
         '--map1', 'PCDJ #50 - Pocket Spring', '--mapper1', '[CSC] variableferret',
         '--map2', 'PCDJ #50 - Null Star', '--mapper2', '[Fenn]Lilly Fenn'],
    # Cup 52's row 3 was written free-form by hand; only its data is compared.
    52: ['--community', '51', '--date', '2026-09-23',
         '--map1', 'PCDJ - Ridge', '--mapper1', 'justMaki',
         '--map2', 'Polymer Composite Derivative Junction', '--mapper2', 'agix'],
    53: ['--community', '52', '--date', '2026-09-30',
         '--map1', 'PCDJ #52 - Bankbanknewnew', '--mapper1', '[CSC] Shadynook',
         '--map2', 'PCDJ #52 - A Little Bit of Bob', '--mapper2', '[CSC] redal'],
}


def block(path, col):
    ws = openpyxl.load_workbook(path).active
    rows = []
    for r in range(2, ws.max_row + 1):
        if r in (3, 4):
            continue
        vals = [ws.cell(r, c).value for c in range(col, col + 4)]
        if r > 5 and not any(v is not None for v in vals):
            break
        rows.append(vals)
    return rows, ws.cell(3, col).value


def new_petite(repo, n, *extra):
    inbox = repo / 'inbox'
    inbox.mkdir(exist_ok=True)
    shutil.copy2(LOGS / f'petite_{n}.log', inbox / 'LogOutput.log')
    return run_script(repo, 'new_petite.py', *CUPS[n], '--log', str(inbox / 'LogOutput.log'), '--no-open', *extra)


@pytest.fixture
def emptied(scratch):
    """The scratch repo as it was after Cup 50: Season 4 removed. Returns (repo, baseline rankings).
    The baseline is the rankings with exactly the replayed cups (51-53), whatever came later."""
    meta_path = scratch / 'cup_meta.json'
    meta = json.loads(meta_path.read_text(encoding='utf-8'))
    meta = {k: v for k, v in meta.items() if int(k.split()[-1]) in CUPS}
    meta_path.write_text(json.dumps(meta, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    rc, out = run_script(scratch, 'petite_ranking.py')
    assert rc == 0, out
    baseline = (scratch / 'petite_rankings.json').read_bytes()
    (scratch / BOOK).rename(scratch / 'original_51-55.xlsx')
    (scratch / 'cup_meta.json').write_text('{}\n', encoding='utf-8')
    return scratch, baseline


def test_replay_51_to_53(emptied):
    repo, baseline = emptied
    for n in (51, 52, 53):
        rc, out = new_petite(repo, n)
        assert rc == 0, out
        assert f'PETITE CUP {n} COMPLETE' in out
        assert '(no livelog)' in out

    for n, col in ((51, 1), (52, 7), (53, 13)):
        got, got_maps = block(repo / BOOK, col)
        want, want_maps = block(repo / 'original_51-55.xlsx', col)
        assert got == want, f'Cup {n} xlsx block differs'
        if n != 52:
            assert got_maps == want_maps

    assert (repo / 'petite_rankings.json').read_bytes() == baseline

    meta = json.loads((repo / 'cup_meta.json').read_text(encoding='utf-8'))
    seeds = json.loads((REPO / 'cup_meta.json').read_text(encoding='utf-8'))
    for label, seed in seeds.items():
        if int(label.split()[-1]) not in CUPS:
            continue
        for key in ('season', 'round', 'xlsx', 'date', 'community', 'excluded'):
            assert meta[label][key] == seed[key], (label, key)
    assert meta['Petite Cup 51']['created_xlsx'] is True
    assert json.loads((repo / 'status.json').read_text(encoding='utf-8'))['as_of_cup'] == 52


def test_guards(scratch):
    """Run against the repo's real latest cup, whichever it is."""
    import datetime
    repo = scratch
    meta = json.loads((repo / 'cup_meta.json').read_text(encoding='utf-8'))
    last = max(meta.values(), key=lambda m: (m['date'], m['community']))
    nxt, last_date = last['community'] + 1, datetime.date.fromisoformat(last['date'])
    log = ['--log', str(LOGS / 'petite_53.log'), '--no-open']

    def args(community, date):
        return ['--community', str(community), '--date', date.isoformat(),
                '--map1', 'A', '--map2', 'B'] + log

    rc, out = run_script(repo, 'new_petite.py', *args(last['community'], last_date))
    assert rc == 1 and 'is already processed' in out

    # A two-week date gap means a cup is missing in between.
    rc, out = run_script(repo, 'new_petite.py', *args(nxt, last_date + datetime.timedelta(days=14)))
    assert rc == 1 and 'not one week' in out

    # A community number that skips one.
    rc, out = run_script(repo, 'new_petite.py', *args(nxt + 1, last_date + datetime.timedelta(days=7)))
    assert rc == 1 and f'does not follow PCDJ #{last["community"]}' in out

    # Nothing was written by the refusals.
    assert json.loads((repo / 'cup_meta.json').read_text(encoding='utf-8')) == meta


def test_reprocess_latest(emptied):
    repo, baseline = emptied
    for n in (51, 52, 53):
        rc, out = new_petite(repo, n)
        assert rc == 0, out
    first = (repo / BOOK).read_bytes()

    rc, out = new_petite(repo, 52, '--reprocess')
    assert rc == 1 and 'only the latest cup' in out

    rc, out = new_petite(repo, 53, '--reprocess')
    assert rc == 0, out
    assert 'already holds the standings after Cup 52' in out
    assert (repo / 'petite_rankings.json').read_bytes() == baseline
    got, _ = block(repo / BOOK, 13)
    want, _ = block(repo / 'original_51-55.xlsx', 13)
    assert got == want
    assert (repo / 'backups' / f'{BOOK}.pre_reprocess_53.xlsx').exists()
