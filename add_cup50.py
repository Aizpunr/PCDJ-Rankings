"""Add Cup 50 (Troll 4) to Petite Cups 46-50.xlsx in cols 25-28.

Source: cup logs/petite_50_reconstructed.json (parsed from local LogOutput.log;
LiveLeaderboardLogger was stale Jun 7 so map names come from old data — see
maps placeholder below; edit row 3 manually once confirmed).

No mappers in the lobby — all 30 lobby participants ranked.
aizpun joined the lobby but DNF'd R1 (effectively didn't compete).
"""
import json
import os
import openpyxl

base = os.path.dirname(os.path.abspath(__file__))
xlsx_path = os.path.join(base, 'Petite Cups 46-50.xlsx')
json_path = os.path.join(base, 'cup logs', 'petite_50_reconstructed.json')

EXCLUDE = set()

with open(json_path, 'r', encoding='utf-8') as f:
    raw = json.load(f)

filtered = [r for r in raw if r['name'] not in EXCLUDE]

new_list = []
prev_pos = None
new_pos = 0
seen = 0
for entry in filtered:
    seen += 1
    if entry['pos'] != prev_pos:
        new_pos = seen
        prev_pos = entry['pos']
    new_list.append({**entry, 'new_pos': new_pos})

print(f'Total entries after mapper exclusion: {len(new_list)}')
for e in new_list:
    t = f"{e['time']:.5f}" if e['time'] is not None else 'DNF'
    safe = e['name'].encode('ascii', 'replace').decode('ascii')
    print(f"  {e['new_pos']:3} {safe:38} {t:>10}  (R{e['round'] or '-'})")

wb = openpyxl.load_workbook(xlsx_path)
ws = wb.active

COL = 25  # cup 46 at 1, 47 at 7, 48 at 13, 49 at 19, 50 at 25 (6-col groups)
ws.cell(row=2, column=COL, value='Petite Cup 50 (Troll 4)')
ws.cell(row=3, column=COL, value='Maps: PCDJ Troll #3 - Deja Vu by [CSC] An Actual g00se + PCDJ Troll 3 - Wavey Woods by [CSC] redal')
ws.cell(row=5, column=COL,   value='Position')
ws.cell(row=5, column=COL+1, value='Name')
ws.cell(row=5, column=COL+2, value='Elim Time')
ws.cell(row=5, column=COL+3, value='Elim Round')

for r in range(6, 60):
    for c in range(COL, COL+4):
        ws.cell(row=r, column=c, value=None)

for i, e in enumerate(new_list):
    r = 6 + i
    ws.cell(row=r, column=COL,   value=e['new_pos'])
    ws.cell(row=r, column=COL+1, value=e['name'])
    if e['time'] is not None:
        ws.cell(row=r, column=COL+2, value=round(e['time'], 5))
    else:
        ws.cell(row=r, column=COL+2, value='DNF')
    if e['round'] is not None and e.get('note') != 'winner':
        ws.cell(row=r, column=COL+3, value=e['round'])

wb.save(xlsx_path)
print(f'\nWrote Cup 50 -> {xlsx_path}')
