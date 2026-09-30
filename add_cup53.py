"""Add Cup 53 (S4 Round 3) to Petite Cups 51-55.xlsx in cols 13-16.

S4 Round 3, 2026-09-30 (community "Petite Cup Du Jour 52"). No discovery rounds
(every round ran 82-84s), so the 14 COTDTracker elim rounds map 1:1 onto plugin
rounds 8-21.

AFTER RUNNING, in petite_ranking.py add to the Season 4 dicts:
  FULL_LOBBY_REPLACEMENTS['Season 4']['Round 3'] = ('Petite Cups 51-55.xlsx', 'Petite Cup 53')
  EVENT_DATES['Season 4']['Round 3'] = '2026-09-30'
Both are required -- the xlsx alone puts data in the file but not in the rankings.
"""
import json
import os
import openpyxl

base = os.path.dirname(os.path.abspath(__file__))
xlsx_path = os.path.join(base, 'Petite Cups 51-55.xlsx')
json_path = os.path.join(base, 'cup logs', 'petite_53_reconstructed.json')

# Mappers were Shadynook and redal (per the cup announcement); neither joined the
# lobby -- all 29 roster entries raced -- so there is nothing to exclude.
EXCLUDE = set()
MAPS = ('Maps: PCDJ #52 - Bankbanknewnew by [CSC] Shadynook'
        ' + PCDJ #52 - A Little Bit of Bob by [CSC] redal')

with open(json_path, 'r', encoding='utf-8') as f:
    raw = json.load(f)

filtered = [r for r in raw if r['name'] not in EXCLUDE]

new_list = []
prev_pos = None
new_pos = 0
seen = 0
for entry in filtered:
    seen += 1
    if entry['pos'] != prev_pos:
        new_pos = seen
        prev_pos = entry['pos']
    new_list.append({**entry, 'new_pos': new_pos})

print(f'Total entries after mapper exclusion: {len(new_list)}')
for e in new_list:
    t = f"{e['time']:.5f}" if e['time'] is not None else 'DNF'
    safe = e['name'].encode('ascii', 'replace').decode('ascii')
    print(f"  {e['new_pos']:3} {safe:38} {t:>10}  (R{e['round'] or '-'})")

wb = openpyxl.load_workbook(xlsx_path)
ws = wb.active

COL = 13  # cup 51 at 1, 52 at 7, 53 at 13, 54 at 19, 55 at 25 (6-col groups)
ws.cell(row=2, column=COL, value='Petite Cup 53')
ws.cell(row=3, column=COL, value=MAPS)
ws.cell(row=5, column=COL,   value='Position')
ws.cell(row=5, column=COL+1, value='Name')
ws.cell(row=5, column=COL+2, value='Elim Time')
ws.cell(row=5, column=COL+3, value='Elim Round')

for r in range(6, 60):
    for c in range(COL, COL+4):
        ws.cell(row=r, column=c).value = None

for i, e in enumerate(new_list):
    r = 6 + i
    ws.cell(row=r, column=COL,   value=e['new_pos'])
    ws.cell(row=r, column=COL+1, value=e['name'])
    if e['time'] is not None:
        ws.cell(row=r, column=COL+2, value=round(e['time'], 5))
    else:
        ws.cell(row=r, column=COL+2, value='DNF')
    if e['round'] is not None and e.get('note') != 'winner':
        ws.cell(row=r, column=COL+3, value=e['round'])

wb.save(xlsx_path)
print(f'\nWrote Cup 53 -> {xlsx_path}')
